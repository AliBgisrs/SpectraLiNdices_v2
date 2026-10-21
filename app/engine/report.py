"""Excel workbook writer (one sheet per statistic + analysis sheets)."""
from __future__ import annotations

import re

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .meta import CREDIT, CREDIT_LONG

HEADER_FILL = PatternFill("solid", fgColor="1F4E3D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SIG_FILL = PatternFill("solid", fgColor="FCE4D6")

VAR_DOC = {
    "Green": "Green reflectance (0-1)", "Red": "Red reflectance (0-1)", "RedEdge": "Red-edge reflectance (0-1)",
    "NIR": "Near-infrared reflectance (0-1)", "Blue": "Blue reflectance (0-1)",
    "Thermal_C": "Surface temperature, all plot pixels (deg C)",
    "Tcanopy_C": "Canopy temperature, vegetation pixels only (deg C)",
    "Tsoil_C": "Soil temperature, bare-soil pixels only (deg C)",
    "PlantHeight_m": "Plant height = DSM - generated DTM, vegetation pixels only (m)",
    "DSM_m": "Digital surface model elevation (m)", "DTM_m": "Generated bare-earth terrain elevation (m)",
}


def _safe(name, used):
    n = re.sub(r"[\[\]\*\?/\\:]", "_", str(name))[:31]
    base, i = n, 1
    while n.lower() in used:
        suf = f"_{i}"
        n = base[: 31 - len(suf)] + suf
        i += 1
    used.add(n.lower())
    return n


def _style(ws, freeze="B2", width=None):
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 50)), 1):
        w = width or min(45, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w


def write_workbook(path, tables, stat_sheets, height_pct, adv_df, thermal_df, thermal_refs, gdd_df, weather,
                   stat_res, pair_res, corr, merged, indices, meta, errors, group_series, predictors=None):
    used = set()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        def put(df, name, index=True, freeze="B2", width=None):
            s = _safe(name, used)
            df.to_excel(xw, sheet_name=s, index=index)
            ws = xw.sheets[s]
            _style(ws, freeze=freeze, width=width)
            return ws

        # README / metadata
        rows = [("Item", "Value")] + [(k, str(v)) for k, v in meta.items()]
        rows += [("", ""), ("Sheet", "Content")]
        rows += [(stat_sheets[k], f"{stat_sheets[k]} of every variable per plot (rows = PLOTID)") for k in stat_sheets]
        rows += [("PixelCount", "Number of pixels used for each variable and plot"),
                 ("PlantHeight_Pctl", "Upper percentiles of plant height (p90/p95/p99 are the usual 'plot height')"),
                 ("Canopy_Structure", "Canopy cover, area, volume, harvestable area, stand count, canopy/plant length"),
                 ("Thermal_Advanced / Thermal_References", "Canopy temperature indicators (CTD, CWSI, Ig, I3, NRCT) and field wet/dry references + QC flags"),
                 ("GDD / Weather_Daily", "Growing degree days per plot and the downloaded daily weather"),
                 ("Stats_Design / Stats_Omnibus", "Test design + omnibus test per variable (p, BH across variables, effect size)"),
                 ("Stats_Letters", "Compact letter display per plot: plots sharing a letter are not significantly different"),
                 ("Stats_PlotSummary / Stats_GroupSummary", "n, mean, SD, SE, 95% CI, median, IQR, letters per plot / group"),
                 ("Stats_Pairs / PW_<var>", "Pairwise post-hoc tests (long format) and adjusted p-value matrices"),
                 ("Corr_*", "Correlation matrices between bands / indices"),
                 ("All_Predictors", "One row per PLOTID with every numeric image-derived value (also written as its own PlotPredictors_*.xlsx)"),
                 ("Integrated_FieldData", "Field data joined with all numeric image-derived plot values"),
                 ("Indices_Reference", "Formulas and references of the vegetation indices"),
                 ("", ""), ("Variable", "Definition")]
        rows += [(k, v) for k, v in VAR_DOC.items()]
        rows += [(ix["abbr"], f"{ix['name']}: {ix['formula']}") for ix in indices]
        ws = put(pd.DataFrame(rows[1:], columns=list(rows[0])), "README", index=False, freeze="A2")
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 110

        for k, sname in stat_sheets.items():
            df = tables[k].copy()
            if group_series is not None:
                df.insert(0, group_series.name, group_series.reindex(df.index).values)
            put(df, sname)
        put(tables["count"], "PixelCount")
        if height_pct is not None and not height_pct.empty:
            put(height_pct, "PlantHeight_Pctl")
        if adv_df is not None and not adv_df.empty:
            put(adv_df, "Canopy_Structure")
        if thermal_df is not None and not thermal_df.empty:
            put(thermal_df, "Thermal_Advanced")
            if thermal_refs:
                put(pd.DataFrame(list(thermal_refs.items()), columns=["Item", "Value"]).astype({"Value": str}),
                    "Thermal_References", index=False, freeze="A2")
        if gdd_df is not None and not gdd_df.empty:
            put(gdd_df, "GDD")
        if weather is not None:
            w = pd.DataFrame({"date": weather["dates"], "Tmax_C": weather["tmax"], "Tmin_C": weather["tmin"]})
            w["Tmean_C"] = (w["Tmax_C"] + w["Tmin_C"]) / 2
            ws = put(w, "Weather_Daily", index=False, freeze="A2")
            ws.cell(1, 6, "Source").font = Font(bold=True)
            ws.cell(2, 6, weather["source"])
            ws.cell(3, 6, weather["url"])

        # statistics
        if stat_res:
            from .stats import DISCLAIMER
            design = [("Alpha", meta.get("Alpha", 0.05)), ("Replicate unit (Mode A)", meta.get("Replicate unit (stats)", "")),
                      ("Mode A", "plot vs plot - " + DISCLAIMER),
                      ("Mode B", "treatment groups: plot means are replicates (CRD one-way, or RCBD additive with a block column)"),
                      ("Post-hoc (Mode A)", "k<=60: Games-Howell (Welch path) or Dunn (KW path) + letters; 60<k<=150: Dunn; "
                                            "k>150: each plot vs rest of field (Mann-Whitney U, BH)"),
                      ("Post-hoc (Mode B)", "Tukey HSD / Games-Howell / Dunn-Holm by assumption checks; only if omnibus p < alpha"),
                      ("Letters", "Piepho (2004) insert-absorb-sweep; 'a' = highest mean; sharing a letter = not different"),
                      ("p_omnibus_BH", "Benjamini-Hochberg across all variables of this run"),
                      ("sig_practical", "significant AND |difference| >= smallest meaningful difference (delta_practical)")]
            put(pd.DataFrame(design, columns=["Item", "Value"]), "Stats_Design", index=False, freeze="A2")
            sdf = pd.DataFrame(list(stat_res.values()))
            front = [c for c in ("Variable", "Mode", "primary_test", "k", "N", "p_primary", "p_omnibus_BH", "significant",
                                 "effect_size", "effect_label", "flags") if c in sdf.columns]
            sdf = sdf[front + [c for c in sdf.columns if c not in front]]
            ws = put(sdf, "Stats_Omnibus", index=False, freeze="C2")
            if "significant" in sdf.columns:
                for i, val in enumerate(sdf["significant"], 2):
                    if val is True or val == 1:
                        for c in ws[i]:
                            c.fill = SIG_FILL
            lett = {v: p["letters"] for v, p in pair_res.items() if p and p.get("letters") and not v.endswith("|groups")}
            if lett:
                ldf = pd.DataFrame(lett)
                ldf.index.name = "PLOTID"
                put(ldf, "Stats_Letters")
            for mode, key in (("A", lambda v: not v.endswith("|groups")), ("B", lambda v: v.endswith("|groups"))):
                rows = []
                for v, p in pair_res.items():
                    if p and key(v) and p.get("summary") is not None:
                        d = p["summary"].copy()
                        d.insert(0, "Variable", v.replace("|groups", ""))
                        d.insert(1, "Posthoc", p.get("posthoc"))
                        rows.append(d)
                if rows:
                    d = pd.concat(rows, ignore_index=True).rename(columns={"Unit": "PLOTID" if mode == "A" else "Group"})
                    if len(d) < 1_000_000:
                        put(d, "Stats_PlotSummary" if mode == "A" else "Stats_GroupSummary", index=False, freeze="C2")
            prs = []
            for v, p in pair_res.items():
                if p and p.get("pairs") is not None:
                    d = p["pairs"].copy()
                    if not v.endswith("|groups") and len(p["names"]) > 60:
                        d = d[d["sig_stat"]]  # many plots: keep the workbook light, list significant pairs only
                    d.insert(0, "Variable", v.replace("|groups", ""))
                    d.insert(1, "Mode", "B" if v.endswith("|groups") else "A")
                    d.insert(2, "Posthoc", p.get("posthoc"))
                    prs.append(d)
            if prs:
                d = pd.concat(prs, ignore_index=True)
                if len(d) < 1_000_000:
                    put(d, "Stats_Pairs", index=False, freeze="D2")
            for v, p in pair_res.items():
                if not p or p.get("P") is None or len(p["names"]) > 60:
                    continue
                pdf = p["P"].copy()
                pdf.index.name = p["posthoc"]
                ws = put(pdf, "PW_" + v.replace("|groups", "_grp"), width=9)
                for row in ws.iter_rows(min_row=2, min_col=2):
                    for c in row:
                        if isinstance(c.value, float):
                            c.number_format = "0.0000"
                            if c.value < float(meta.get("Alpha", 0.05)):
                                c.fill = SIG_FILL

        if corr:
            for key, name in (("across", "Corr_AcrossPlots"), ("across_p", "Corr_AcrossPlots_p"),
                              ("within_mean", "Corr_WithinPlot_mean")):
                if corr.get(key) is not None:
                    put(corr[key].round(4), name, width=11)
            if corr.get("within"):
                rows = []
                for pid, m in corr["within"].items():
                    cols = list(m.columns)
                    for i in range(len(cols)):
                        for j in range(i + 1, len(cols)):
                            rows.append((pid, cols[i], cols[j], m.iat[i, j]))
                ldf = pd.DataFrame(rows, columns=["PLOTID", "Var1", "Var2", "r"])
                if len(ldf) < 1_000_000:
                    put(ldf, "Corr_WithinPlot_long", index=False, freeze="A2")

        if predictors is not None:
            put(predictors, "All_Predictors")
        if merged is not None:
            put(merged, "Integrated_FieldData", index=False)

        idf = pd.DataFrame([{k: ix.get(k, "") for k in ("abbr", "name", "formula", "sensitivity", "reference", "doi")}
                            for ix in indices])
        put(idf, "Indices_Reference", index=False, freeze="A2")
        if errors:
            put(pd.DataFrame(errors, columns=["PLOTID", "Problem"]), "Failed_Plots", index=False, freeze="A2")
    return path


PRED_DOC = {
    "min": "minimum", "mean": "mean", "median": "median", "max": "maximum", "std": "standard deviation (ddof=1)",
    "range": "max - min", "count": "number of pixels used",
}


def write_predictors(path, predictors, dictionary, meta):
    """Stand-alone workbook: one row per PLOTID, every numeric image-derived predictor as a column."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        predictors.to_excel(xw, sheet_name="Predictors")
        _style(xw.sheets["Predictors"])
        dictionary.to_excel(xw, sheet_name="Dictionary", index=False)
        _style(xw.sheets["Dictionary"], freeze="A2")
        xw.sheets["Dictionary"].column_dimensions["C"].width = 90
        info = pd.DataFrame([("Generated by", CREDIT_LONG)] + [(k, str(v)) for k, v in meta.items()
                                                                if k in ("Input folder", "Boundary file", "Plot ID column",
                                                                         "Plots analysed", "Image date/time", "Created")],
                            columns=["Item", "Value"])
        info.to_excel(xw, sheet_name="About", index=False)
        _style(xw.sheets["About"], freeze="A2")
        xw.sheets["About"].column_dimensions["B"].width = 90
    return path
